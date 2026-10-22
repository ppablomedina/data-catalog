#!/usr/bin/env python3
"""Convierte el inventario XLSX de SAGULPA en datos estáticos del portal."""

from __future__ import annotations

import json
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


AREA_ALIASES = {
    "financiero": "financiero",
    "sistemas": "sistemas",
    "oficina del dato": "oficina-dato",
    "ser": "ser",
    "juridico": "juridico",
    "tecnico": "tecnico",
    "seguridad": "seguridad-salud-calidad",
    "seguridad, salud y calidad": "seguridad-salud-calidad",
    "atencion": "atencion",
    "financiacion": "financiacion-externa",
    "financiacion externa": "financiacion-externa",
    "sostenibilidad y energia": "sostenibilidad-energia",
    "transversal": "transversal",
    "aparcamientos": "aparcamientos",
    "centro de incidencias operativas": "centro-incidencias-operativas",
    "moxsi": "moxsi",
    "movilidad": "movilidad-trafico",
    "movilidad y trafico": "movilidad-trafico",
    "personas y talento": "personas-talento",
}

GOVERNANCE_LABELS = {
    "ownerArea": "el área propietaria",
    "data": "los datos aportados",
    "useCases": "los casos de uso",
    "access": "el formato o vía de acceso",
    "frequency": "la frecuencia",
    "history": "el histórico",
    "quality": "la calidad",
    "personalData": "los datos personales",
    "rgpdCategory": "la categoría RGPD",
    "classification": "la clasificación",
    "criticality": "la criticidad",
    "blocker": "los bloqueos o pendientes",
}


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        return re.sub(r"\s+", " ", value).strip()
    return value


def normalized(value: str) -> str:
    return "".join(
        character
        for character in unicodedata.normalize("NFD", value.casefold())
        if unicodedata.category(character) != "Mn"
    )


def slug(value: str) -> str:
    text = normalized(value)
    text = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    return text[:70]


def area_ids(owner_area: str) -> list[str]:
    parts = [clean(part) for part in owner_area.split("/")]
    result = []
    for part in parts:
        identifier = AREA_ALIASES.get(normalized(part))
        if not identifier:
            raise ValueError(f"Área sin correspondencia: {part!r}")
        if identifier not in result:
            result.append(identifier)
    return result


def icon_for(name: str, access: str) -> str:
    text = normalized(f"{name} {access}")
    if "api" in text:
        return "code"
    if any(word in text for word in ("telemetria", "sensor", "aforador", "monitor")):
        return "chart"
    if any(word in text for word in ("matricula", "trafico", "ubicacion", "zona")):
        return "pin"
    if any(word in text for word in ("bateria", "vehiculo", "parquimetro")):
        return "battery"
    if any(word in text for word in ("ticket", "incidencia", "tarea", "glpi", "jira")):
        return "tasks"
    if any(
        word in text
        for word in (
            "document",
            "pdf",
            "pliego",
            "normativa",
            "proyecto",
            "expediente",
            "acta",
        )
    ):
        return "folder"
    return "database"


def is_incomplete(value) -> bool:
    if value is None or clean(value) == "":
        return True
    text = normalized(str(value))
    return "(pendiente)" in text


def access_status(access: str) -> str:
    text = normalized(access)
    if "api" in text and not any(
        marker in text for marker in ("pendiente", "confirmar", "validar")
    ):
        return "available"
    if any(marker in text for marker in ("pendiente", "confirmar", "validar")):
        return "pending"
    return "direct"


def load_inventory(workbook_path: Path):
    workbook = load_workbook(workbook_path, data_only=True)
    sheet = workbook["Inventario de fuentes"]
    records = []

    for values in sheet.iter_rows(min_row=4, max_col=14, values_only=True):
        if values[0] is None:
            continue
        (
            inventory_id,
            name,
            owner_area,
            data,
            use_cases,
            access,
            frequency,
            history,
            quality,
            personal_data,
            rgpd_category,
            classification,
            criticality,
            blocker,
        ) = [clean(value) for value in values]
        governance = {
            "inventoryId": inventory_id,
            "ownerArea": owner_area,
            "data": data,
            "useCases": use_cases,
            "access": access,
            "frequency": frequency,
            "history": history,
            "quality": quality,
            "personalData": personal_data,
            "rgpdCategory": rgpd_category,
            "classification": classification,
            "criticality": criticality,
            "blocker": blocker,
        }
        completeness_fields = [
            "ownerArea",
            "data",
            "useCases",
            "access",
            "frequency",
            "history",
            "quality",
            "personalData",
            "rgpdCategory",
            "classification",
            "criticality",
            "blocker",
        ]
        missing = [
            field for field in completeness_fields if is_incomplete(governance[field])
        ]
        pending = []
        if blocker:
            pending.append(blocker)
        else:
            pending.append("Confirmar si existen bloqueos o tareas pendientes.")
        pending.extend(
            f"Completar y validar {GOVERNANCE_LABELS[field]}."
            for field in missing
            if field != "blocker"
        )
        total = len(completeness_fields)
        filled = total - len(missing)
        areas = area_ids(owner_area)
        source_icon = icon_for(name, access or "")
        records.append(
            {
                "id": f"inventario-{inventory_id}-{slug(name)}",
                "inventoryId": inventory_id,
                "departmentId": areas[0],
                "departmentIds": areas,
                "serviceIds": (
                    ["sitycleta-mecanica", "sitycleta-electrica", "sityneta"]
                    if "moxsi" in areas
                    else []
                ),
                "name": name,
                "initials": str(inventory_id),
                "color": "blue",
                "category": "Inventario IA",
                "logoIcon": source_icon,
                "description": data or "Información pendiente de documentar.",
                "contentTitle": "Información inventariada",
                "contentDescription": (
                    "Datos identificados en las entrevistas por área para los "
                    "casos de uso priorizados."
                ),
                "itemLabel": "registro",
                "datasets": [
                    {
                        "name": name,
                        "description": data or "Información pendiente de documentar.",
                        "icon": source_icon,
                        "status": "Inventariado",
                    }
                ],
                "api": {
                    "status": access_status(access or ""),
                    "url": None,
                    "note": access or "Vía de acceso pendiente de documentar.",
                },
                "website": None,
                "repository": None,
                "storage": None,
                "owner": owner_area,
                "pending": pending,
                "completeness": {
                    "filled": filled,
                    "total": total,
                    "percent": round(filled / total * 100),
                    "missing": missing,
                },
                "governance": governance,
            }
        )

    use_case_sheet = workbook["Mapa caso ↔ fuente"]
    use_cases = []
    for row in use_case_sheet.iter_rows(min_row=4, max_col=5, values_only=True):
        if not row[0]:
            continue
        use_cases.append(
            {
                "name": clean(row[0]),
                "area": clean(row[1]),
                "sources": clean(row[2]),
                "sourceCount": clean(row[3]),
                "minimumDataAvailable": clean(row[4]),
            }
        )

    gaps_sheet = workbook["Gaps y peticiones"]
    gaps = []
    for row in gaps_sheet.iter_rows(min_row=4, max_col=6, values_only=True):
        if row[0] is None:
            continue
        gaps.append(
            {
                "id": clean(row[0]),
                "question": clean(row[1]),
                "directedTo": clean(row[2]),
                "source": clean(row[3]),
                "priority": clean(row[4]),
                "response": clean(row[5]),
            }
        )

    criteria_sheet = workbook["Criterios y leyenda"]

    def criteria_rows(start, end):
        return [
            {"name": clean(criteria_sheet.cell(row, 1).value), "description": clean(criteria_sheet.cell(row, 2).value)}
            for row in range(start, end + 1)
        ]

    metadata = {
        "title": clean(sheet.cell(1, 1).value),
        "scope": clean(sheet.cell(2, 1).value),
        "useCases": use_cases,
        "gaps": gaps,
        "criteria": {
            "qualityLegend": criteria_rows(4, 7),
            "qualityCriteria": criteria_rows(10, 14),
            "sensitivity": criteria_rows(17, 20),
            "minimumViable": criteria_rows(23, 26),
        },
    }
    return records, metadata


def main():
    if len(sys.argv) != 3:
        raise SystemExit("Uso: import_inventory.py <entrada.xlsx> <salida.js>")
    input_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    records, metadata = load_inventory(input_path)
    output = (
        "/* Generado por scripts/import_inventory.py. No editar manualmente. */\n"
        f"window.DATA_INVENTORY = {json.dumps(records, ensure_ascii=False, separators=(',', ':'))};\n"
        f"window.INVENTORY_META = {json.dumps(metadata, ensure_ascii=False, separators=(',', ':'))};\n"
    )
    output_path.write_text(output, encoding="utf-8")
    print(f"{len(records)} fuentes escritas en {output_path}")


if __name__ == "__main__":
    main()
